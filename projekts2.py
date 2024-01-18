from openpyxl import Workbook, load_workbook

filiale=input("Izvēlieties filiāli (Rīga, Mārupe, Ādaži): ")
diena=input("Izvēlieties dienu: ")

wb=load_workbook('klasu grafiks 2023_2024.xlsx')
ws1=wb['Rīga']
ws2=wb['Mārupe']
ws3=wb['Ādaži']

max_row1=ws1.max_row
max_row2=ws2.max_row
max_row3=ws3.max_row

brivais_laiks_pirmdiena=False
brivais_laiks_otrdiena=False
brivais_laiks_tresdiena=False
brivais_laiks_ceturtdiena=False
brivais_laiks_piektdiena=False

if filiale == 'Rīga':
    for row in range(3,13):
        brivie_laiki_pirmdiena = ws1['d' + str(row)].value
        brivs1 = ws1['b' + str(row)].value
        brivie_laiki_otrdiena = ws1['i' + str(row)].value
        brivs2 = ws1['g' + str(row)].value
        brivie_laiki_tresdiena = ws1['n' + str(row)].value
        brivs3 = ws1['l' + str(row)].value
        brivie_laiki_ceturtdiena = ws1['s' + str(row)].value
        brivs4 = ws1['q' + str(row)].value
        brivie_laiki_piektdiena = ws1['x' + str(row)].value
        brivs5 = ws1['v' + str(row)].value

        if diena == 'Pirmdiena' and brivs1 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_pirmdiena))
            brivais_laiks_pirmdiena = True
            break
        elif diena == 'Otrdiena' and brivs2 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_otrdiena))
            brivais_laiks_otrdiena = True
            break
        elif diena == 'Trešdiena' and brivs3 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_tresdiena))
            brivais_laiks_tresdiena = True
            break
        elif diena == 'Ceturtdiena' and brivs4 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_ceturtdiena))
            brivais_laiks_ceturtdiena = True
            break
        elif diena == 'Piektdiena' and brivs5 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_piektdiena))
            brivais_laiks_piektdiena = True
            break
    if not brivais_laiks_pirmdiena and not brivais_laiks_otrdiena and not brivais_laiks_tresdiena and not brivais_laiks_ceturtdiena and not brivais_laiks_piektdiena:
        print("Nav brīvu laiku")



if filiale == 'Mārupe':
    for row in range(3,13):
        brivie_laiki_pirmdiena = ws2['d' + str(row)].value
        brivs1 = ws2['b' + str(row)].value
        brivie_laiki_otrdiena = ws2['i' + str(row)].value
        brivs2 = ws2['g' + str(row)].value
        brivie_laiki_tresdiena = ws2['n' + str(row)].value
        brivs3 = ws2['l' + str(row)].value
        brivie_laiki_ceturtdiena = ws2['s' + str(row)].value
        brivs4 = ws2['q' + str(row)].value
        brivie_laiki_piektdiena = ws2['x' + str(row)].value
        brivs5 = ws2['v' + str(row)].value

        if diena == 'Pirmdiena' and brivs1 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_pirmdiena))
            brivais_laiks_pirmdiena = True
            break
        elif diena == 'Otrdiena' and brivs2 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_otrdiena))
            brivs_laiks_otrdiena = True
            break
        elif diena == 'Trešdiena' and brivs3 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_tresdiena))
            brivs_laiks_tresdiena = True
            break
        elif diena == 'Ceturtdiena' and brivs4 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_ceturtdiena))
            brivais_laiks_ceturtdiena = True
            break
        elif diena == 'Piektdiena' and brivs5 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_piektdiena))
            brivs_laiks_piektdiena = True
            break

    if not brivais_laiks_pirmdiena and not brivais_laiks_otrdiena and not brivais_laiks_tresdiena and not brivais_laiks_ceturtdiena and not brivais_laiks_piektdiena:
        print("Nav brīvu laiku")

if filiale == 'Ādaži':
    for row in range(3,13):
        brivie_laiki_pirmdiena = ws3['d' + str(row)].value
        brivs1 = ws3['b' + str(row)].value
        brivie_laiki_otrdiena = ws3['i' + str(row)].value
        brivs2 = ws3['g' + str(row)].value
        brivie_laiki_tresdiena = ws3['n' + str(row)].value
        brivs3 = ws3['l' + str(row)].value
        brivie_laiki_ceturtdiena = ws3['s' + str(row)].value
        brivs4 = ws3['q' + str(row)].value
        brivie_laiki_piektdiena = ws3['x' + str(row)].value
        brivs5 = ws3['v' + str(row)].value

        if diena == 'Pirmdiena' and brivs1 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_pirmdiena))
            brivais_laiks_pirmdiena = True
            break
        elif diena == 'Otrdiena' and brivs2 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_otrdiena))
            brivais_laiks_otrdiena = True
            break
        elif diena == 'Trešdiena' and brivs3 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_tresdiena))
            brivais_laiks_tresdiena = True
            break
        elif diena == 'Ceturtdiena' and brivs4 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_ceturtdiena))
            brivais_laiks_ceturtdiena=True
            break
        elif diena == 'Piektdiena' and brivs5 == None:
            print("Pieejams brīvs laiks - " + str(brivie_laiki_piektdiena))
            brivais_laiks_piektdiena=True
            break

    if not brivais_laiks_pirmdiena and not brivais_laiks_otrdiena and not brivais_laiks_tresdiena and not brivais_laiks_ceturtdiena and not brivais_laiks_piektdiena:
        print("Nav brīvu laiku")